#%%
"""Rebuttal analysis (VShq: "no statistical significance"): vanilla backbone vs. Ours in Table 1.

Reads the per-seed runs of the first three sheets of aaai27.xlsx (table1-<dataset>), recomputes
every mean/std directly from the raw per-seed rows, flags blocks whose spreadsheet AVERAGE/STDEV
range does not cover exactly that block's own rows, and tests Ours > vanilla per cell:
  - Welch's t-test over the 4 seeds (one-sided), Holm-corrected over all cells
  - Mann-Whitney U (one-sided) as a non-parametric check (n=4 vs 4 -> min p = 1/70 = 0.014)

Usage:
  python rebuttal_significance.py --xlsx ~/Desktop/aaai27.xlsx
"""
import re
import argparse

import numpy as np
import openpyxl
from scipy import stats

SHEETS = ["table1-micro_video", "table1-ml-1m", "table1-kuairand"]
MODELS = ["mf", "grurec", "sasrec", "tisasrec", "fearec", "bsarec"]
METRICS = {"recall@10": 5, "ndcg@10": 6}  # 1-based columns E, F


def parse():
    p = argparse.ArgumentParser()
    p.add_argument("--xlsx", default="/Users/wonhyung64/Desktop/aaai27.xlsx")
    p.add_argument("--alpha", type=float, default=0.05)
    return p.parse_args()


def load_blocks(ws_f, ws_v):
    """Return {(block, model): {"rows": [...], "formula_row": r}} for block in {vanilla, ours}."""
    blocks, block_idx = {}, -1
    for r in range(1, ws_v.max_row + 1):
        if ws_v.cell(r, 1).value == "Name":
            block_idx += 1
            continue
        if block_idx > 1:  # third block = debiasing baselines, not part of this test
            break
        model = ws_v.cell(r, 7).value
        if model is None or ws_v.cell(r, 5).value is None:
            continue
        key = ("vanilla" if block_idx == 0 else "ours", model)
        b = blocks.setdefault(key, {"rows": [], "formula_row": None})
        b["rows"].append(r)
        if ws_f.cell(r, 13).value is not None:
            b["formula_row"] = r
    return blocks


def formula_range(formula):
    m = re.search(r"\$E(\d+):\$E(\d+)", formula or "")
    return (int(m.group(1)), int(m.group(2))) if m else None


def holm(pvals):
    order = np.argsort(pvals)
    adj = np.empty(len(pvals))
    running = 0.0
    for rank, i in enumerate(order):
        running = max(running, (len(pvals) - rank) * pvals[i])
        adj[i] = min(1.0, running)
    return adj


def main():
    args = parse()
    wb_f = openpyxl.load_workbook(args.xlsx)                  # formulas
    wb_v = openpyxl.load_workbook(args.xlsx, data_only=True)  # cached values

    cells, audit = [], []
    for sheet in SHEETS:
        ds = sheet.replace("table1-", "")
        blocks = load_blocks(wb_f[sheet], wb_v[sheet])
        for (kind, model), b in blocks.items():
            own = (min(b["rows"]), max(b["rows"]))
            fr = b["formula_row"]
            mean_rng = formula_range(wb_f[sheet].cell(fr, 13).value) if fr else None
            std_rng = formula_range(wb_f[sheet].cell(fr, 15).value) if fr else None
            if mean_rng != own or std_rng != own:
                vals = [wb_v[sheet].cell(r, 5).value for r in b["rows"]]
                audit.append((ds, kind, model, own, mean_rng, std_rng,
                              wb_v[sheet].cell(fr, 13).value, np.mean(vals)))
        for model in MODELS:
            if ("vanilla", model) not in blocks or ("ours", model) not in blocks:
                continue
            row = {"dataset": ds, "model": model}
            for met, col in METRICS.items():
                v = np.array([wb_v[sheet].cell(r, col).value for r in blocks[("vanilla", model)]["rows"]], float)
                o = np.array([wb_v[sheet].cell(r, col).value for r in blocks[("ours", model)]["rows"]], float)
                t = stats.ttest_ind(o, v, equal_var=False, alternative="greater")
                u = stats.mannwhitneyu(o, v, alternative="greater")
                row[met] = {"v": v, "o": o, "p_welch": t.pvalue, "p_mwu": u.pvalue,
                            "t_two_sided": stats.ttest_ind(o, v, equal_var=False).pvalue}
            cells.append(row)

    print("## Spreadsheet audit: AVERAGE/STDEV ranges that are not the block's own rows")
    print("| dataset | block | model | own rows | AVERAGE range | STDEV range | sheet mean | raw mean |")
    print("|---|---|---|---|---|---|---|---|")
    for ds, kind, model, own, mr, sr, sm, rm in audit:
        print(f"| {ds} | {kind} | {model} | E{own[0]}:E{own[1]} | {mr and f'E{mr[0]}:E{mr[1]}'} | "
              f"{sr and f'E{sr[0]}:E{sr[1]}'} | {sm:.4f} | {rm:.4f} |")

    for met in METRICS:
        p_adj = holm(np.array([c[met]["p_welch"] for c in cells]))
        print(f"\n## {met}: Ours vs. vanilla, recomputed from raw per-seed runs (sample std, n=4 each)")
        print("Welch one-sided p (Ours > vanilla), Holm-adjusted over all cells; MWU = Mann-Whitney U one-sided")
        print("| dataset | model | Vanilla | Ours | rel. impr. | Welch p | Holm p | MWU p | sig. |")
        print("|---|---|---|---|---|---|---|---|---|")
        for c, pa in zip(cells, p_adj):
            m = c[met]
            v, o = m["v"], m["o"]
            impr = (o.mean() - v.mean()) / v.mean() * 100
            if o.mean() < v.mean():
                sig = "worse (two-sided p={:.1e})".format(m["t_two_sided"])
            else:
                sig = "yes" if pa < args.alpha else "no"
            print(f"| {c['dataset']} | {c['model']} | {v.mean():.4f}±{v.std(ddof=1):.4f} | "
                  f"{o.mean():.4f}±{o.std(ddof=1):.4f} | {impr:+.1f}% | {m['p_welch']:.1e} | {pa:.1e} | "
                  f"{m['p_mwu']:.3f} | {sig} |")


if __name__ == "__main__":
    main()
